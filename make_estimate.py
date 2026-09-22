# -*- coding: utf-8 -*-
"""
UNION ONE 견적서 생성기  ((주)유니온원)
사용법:  python make_estimate.py input.json [--out 출력폴더] [--assets 자산폴더]

입력 JSON 을 읽어 항상 같은 양식의  견적서.xlsx / 견적서.pdf 를 만든다.
계산(집계·간접비·인지세·절사·목표금액 맞춤)은 전부 이 파일이 담당한다.
"""
import sys, os, json, math, argparse, datetime, shutil, subprocess, urllib.request, re

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
                                Table, TableStyle, NextPageTemplate, PageBreak, Flowable, KeepTogether)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
from PIL import Image as PILImage

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

# ───────────────────────────── 고정 정보 ─────────────────────────────
COMPANY = {
    "name": "(주)유니온원",
    "reg_no": "210-88-03747",
    "ceo": "김정훈",
    "address": "대구광역시 동구 동화천로77길 46, 3층",
    "biz_type": "건설업",
    "biz_item": "철거 및 리모델링",
    "tel": "1551-8757",
    "bank": "KB국민은행 675001-04-342392 (예금주 : (주)유니온원)",
}
DEFAULT_CONDITIONS = [
    "본 견적은 견적일로부터 30일간 유효하며, 이후 자재비·인건비·폐기물 처리비 등의 변동 시 재견적할 수 있습니다.",
    "철거 중 석면 또는 석면 의심 자재 발견 시 작업을 중지하며, 석면조사·해체·처리 비용은 별도 협의합니다.",
    "발주처 요청에 따른 야간·휴일 작업 시 추가 노무비 및 장비비가 발생하며, 노무비는 30% 할증 적용됩니다.",
    "철거 범위 내 기본 보양은 포함하며, 철거 범위 외 시설물 및 추가 요청 보양은 별도 비용으로 적용됩니다.",
    "대금 지급조건은 계약금 50% / 잔금 50%로 하며, 잔금은 공사 완료 후 당일 지급합니다.",
    "본 견적서에 명시되지 않은 추가·변경사항은 공사 전 상호 협의 후 금액 및 작업범위를 결정합니다.",
]
ASSET_URLS = [
    "https://raw.githubusercontent.com/ncore8868/auto_logo_n_dojang/main/{}",
    "https://github.com/ncore8868/auto_logo_n_dojang/raw/main/{}",
]
# 로고·도장 파일명 (옛 로고·도장이 섞여 쓰이지 않도록 전용 이름 사용). 도장 파일이 없으면 서명란은 (인)만 표시
LOGO_FILE = "uo_logo.png"
STAMP_FILE = "uo_dojang.png"
ASSET_DIRS = ["./assets", "/mnt/project", "/mnt/knowledge", "/mnt/user-data/uploads", os.path.dirname(os.path.abspath(__file__)), "."]

# 롯데(대기업) 간접공사비 기준 (2026)
RATE_SANJAE = 0.0356      # 산재보험료 : 노무비
RATE_GOYONG = 0.0101      # 고용보험료 : 노무비
RATE_SAFETY = 0.0311      # 안전관리비 : 직접공사비 (2천만원 이상 공사)
SAFETY_MIN_DIRECT = 20_000_000
RATE_MISC_MAX = 0.10      # 공과잡비 : 직접공사비의 10% 이내
# 인지세 (파트너사 견적서 금액, VAT 제외, 양사 5:5 분담분)
STAMP_TABLE = [(10_000_000, 0), (30_000_000, 9_091), (50_000_000, 18_182),
               (100_000_000, 31_818), (1_000_000_000, 68_182), (float("inf"), 159_091)]

NAVY = colors.HexColor("#1B2A41")
GOLD = colors.HexColor("#B08D57")
NAVY_LIGHT = colors.HexColor("#EEF1F5")
DARK = colors.HexColor("#1F1F1F")
GRAY_TXT = colors.HexColor("#666666")
GRAY_FILL = colors.HexColor("#F3F3F3")
GRAY_FILL2 = colors.HexColor("#E6E6E6")
LINE = colors.HexColor("#BFBFBF")
LINE_DARK = colors.HexColor("#555555")


# ───────────────────────────── 유틸 ─────────────────────────────
def won(n):
    if n is None:
        return ""
    n = int(round(n))
    return "-" if n == 0 else f"{n:,}"


def fmt_qty(q):
    if q is None:
        return ""
    if abs(q - round(q)) < 1e-9:
        return f"{int(round(q)):,}"
    return f"{q:,.2f}".rstrip("0").rstrip(".")


def stamp_tax(amount):
    for limit, tax in STAMP_TABLE:
        if amount <= limit:
            return tax
    return STAMP_TABLE[-1][1]


def korean_money(n):
    """58300000 -> 오천팔백삼십만"""
    n = int(n)
    if n == 0:
        return "영"
    digits = "영일이삼사오육칠팔구"
    small = ["", "십", "백", "천"]
    big = ["", "만", "억", "조", "경"]
    out, unit = [], 0
    while n > 0:
        chunk = n % 10000
        n //= 10000
        if chunk:
            s = ""
            for i in range(3, -1, -1):
                d = (chunk // 10 ** i) % 10
                if d:
                    s += digits[d] + small[i]
            out.append(s + big[unit])
        unit += 1
    return "".join(reversed(out))


def find_asset(name, extra_dir=None):
    """자산 파일 탐색: 지정폴더 → ./assets → 프로젝트 파일 → 업로드 → 스크립트 폴더 → 깃허브"""
    dirs = ([extra_dir] if extra_dir else []) + ASSET_DIRS
    for d in dirs:
        if not d or not os.path.isdir(d):
            continue
        p = os.path.join(d, name)
        if os.path.isfile(p):
            return p
        for root, _dirs, files in os.walk(d):
            if root.count(os.sep) - d.count(os.sep) > 3:
                continue
            if name in files:
                return os.path.join(root, name)
    os.makedirs("./assets", exist_ok=True)
    for tmpl in ASSET_URLS:
        try:
            r = urllib.request.urlopen(tmpl.format(name), timeout=20)
            body = r.read()
            ctype = r.headers.get("content-type", "")
            if r.status == 200 and "text/html" not in ctype and len(body) > 1000 and not body.lstrip().startswith(b"<"):
                dst = os.path.join("./assets", name)
                with open(dst, "wb") as fp:
                    fp.write(body)
                return dst
        except Exception:
            pass
    return None


def _convert_noto_to_ttf():
    """폰트 자산이 없을 때: 시스템 Noto Sans CJK(CFF) 를 한글 부분집합 TrueType 으로 변환 (약 30초)"""
    try:
        from fontTools.ttLib import TTFont as FT, newTable
        from fontTools.subset import Subsetter, Options
        from fontTools.pens.cu2quPen import Cu2QuPen
        from fontTools.pens.ttGlyphPen import TTGlyphPen
    except Exception:
        return None
    srcs = {"Regular": "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", "Bold": "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"}
    if not all(os.path.isfile(v) for v in srcs.values()):
        return None
    uni = set()
    for a_, b_ in [(0x20, 0x7E), (0xA0, 0xFF), (0x2000, 0x206F), (0x20A9, 0x20A9), (0x2100, 0x214F), (0x2190, 0x21FF), (0x2460, 0x24FF),
                   (0x25A0, 0x25FF), (0x2600, 0x26FF), (0x3000, 0x303F), (0x3131, 0x318E), (0x3200, 0x32FF), (0x3380, 0x33FF),
                   (0xAC00, 0xD7A3), (0xFF01, 0xFF5E)]:
        uni.update(range(a_, b_ + 1))
    uni.update(ord(c) for c in "必外內上下大中小別無有前後新舊式計合見積書主副日月火水木金土年時分本社株式會社建設工事")
    os.makedirs("./assets", exist_ok=True)
    out = {}
    for style, src in srcs.items():
        f = FT(src, fontNumber=1)
        opt = Options(); opt.layout_features = []; opt.notdef_outline = True
        opt.drop_tables += ["GSUB", "GPOS", "GDEF", "BASE", "vhea", "vmtx", "VORG"]
        sub = Subsetter(opt); sub.populate(unicodes=sorted(uni)); sub.subset(f)
        go = f.getGlyphOrder(); gs = f.getGlyphSet()
        f["loca"] = newTable("loca"); f["glyf"] = glyf = newTable("glyf"); glyf.glyphOrder = go
        quad = {}
        for gn in go:
            pen = TTGlyphPen(gs); gs[gn].draw(Cu2QuPen(pen, 1.0, reverse_direction=True)); quad[gn] = pen.glyph()
        glyf.glyphs = quad
        del f["CFF "]
        for t in ("VORG", "GSUB", "GPOS", "GDEF", "BASE", "vhea", "vmtx"):
            if t in f:
                del f[t]
        glyf.compile(f)
        hm = f["hmtx"]
        for gn, g in glyf.glyphs.items():
            if hasattr(g, "xMin"):
                hm[gn] = (hm[gn][0], g.xMin)
        mx = newTable("maxp"); mx.tableVersion = 0x00010000
        for k in ("maxZones", "maxTwilightPoints", "maxStorage", "maxFunctionDefs", "maxInstructionDefs", "maxStackElements",
                  "maxSizeOfInstructions", "maxComponentElements", "maxComponentDepth", "maxPoints", "maxContours",
                  "maxCompositePoints", "maxCompositeContours"):
            setattr(mx, k, 0)
        mx.maxZones = 1; f["maxp"] = mx; mx.recalc(f)
        f["post"].formatType = 3.0
        f.sfntVersion = "\x00\x01\x00\x00"
        dst = f"./assets/NotoSansKR-{style}.ttf"
        f.save(dst); out[style] = dst
    return out


def find_font(extra_dir=None):
    """reportlab 은 TrueType 만 읽는다 (시스템 Noto CJK ttc 는 CFF 라서 불가) → TTF 자산 사용, 없으면 변환"""
    for name_r, name_b in [("NotoSansKR-Regular.ttf", "NotoSansKR-Bold.ttf"), ("NanumGothic.ttf", "NanumGothicBold.ttf")]:
        r, b = find_asset(name_r, extra_dir), find_asset(name_b, extra_dir)
        if r and b:
            return r, b, 0
    conv = _convert_noto_to_ttf()
    if conv:
        return conv["Regular"], conv["Bold"], 0
    raise SystemExit("한글 폰트를 찾지 못했습니다. NotoSansKR-Regular.ttf / NotoSansKR-Bold.ttf 를 assets 폴더(또는 깃허브 저장소)에 넣어 주세요.")


# ───────────────────────────── 계산 ─────────────────────────────
def nice(v):
    """단가를 보기 좋은 단위로 반올림"""
    if v <= 0:
        return 0
    step = 5000 if v >= 100000 else 1000 if v >= 20000 else 500 if v >= 5000 else 100 if v >= 1000 else 10
    return max(step, int(round(v / step)) * step)


def _sum_items(sections):
    direct_mat = direct_lab = 0
    for i, sec in enumerate(sections, 1):
        m = l = 0
        for it in sec["items"]:
            qty = float(it.get("qty", 1))
            mat = int(it.get("mat", 0) or 0)
            lab = int(it.get("lab", 0) or 0)
            it["_qty"], it["_mat"], it["_lab"] = qty, mat, lab
            it["_mat_amt"] = int(round(mat * qty))
            it["_lab_amt"] = int(round(lab * qty))
            it["_sum"] = mat + lab
            it["_sum_amt"] = it["_mat_amt"] + it["_lab_amt"]
            m += it["_mat_amt"]
            l += it["_lab_amt"]
        sec["_no"] = i
        sec["_mat"], sec["_lab"], sec["_sum"] = m, l, m + l
        direct_mat += m
        direct_lab += l
    return direct_mat, direct_lab


def _scale(sections, k):
    for sec in sections:
        for it in sec["items"]:
            it["mat"] = nice(int(it.get("mat", 0) or 0) * k)
            it["lab"] = nice(int(it.get("lab", 0) or 0) * k)


def _fit_item(sections, name=None):
    """잔차를 흡수할 항목: 지정 이름 > 수량 1 인 '식' 항목 중 마지막 > 마지막 항목"""
    cand = None
    for sec in sections:
        for it in sec["items"]:
            if name and it.get("name") == name:
                return it
            if it.get("unit") == "식" and float(it.get("qty", 1)) == 1:
                cand = it
    if cand:
        return cand
    return sections[-1]["items"][-1]


def _lotte_indirect(direct, lab):
    sanjae = int(round(lab * RATE_SANJAE))
    goyong = int(round(lab * RATE_GOYONG))
    safety = int(round(direct * RATE_SAFETY)) if direct >= SAFETY_MIN_DIRECT else 0
    return sanjae, goyong, safety


def compute(data):
    mode = data.get("mode", "general")
    proj = data["project"]
    sections = data["sections"]
    target = proj.get("target_total")
    auto_fit = proj.get("auto_fit", True)
    scale_k = 1.0
    fit_note = ""

    direct_mat, direct_lab = _sum_items(sections)
    direct = direct_mat + direct_lab

    # ── 목표금액 자동 맞춤 (단가 비례조정 → 잔차 흡수)
    already_ok = False
    if target and direct > 0:
        if mode == "lotte":
            sj, gy, sf = _lotte_indirect(direct, direct_lab)
            m0 = target - (direct + sj + gy + sf) - stamp_tax(target)
            already_ok = 0 <= m0 <= direct * RATE_MISC_MAX
        else:
            already_ok = target <= direct < target + 10000
    if target and auto_fit and direct > 0 and not already_ok:
        if mode == "lotte":
            ratio = direct_lab / direct
            for _ in range(3):
                tax = stamp_tax(target)
                need = (target - tax) / (1 + RATE_SAFETY + 0.08 + (RATE_SANJAE + RATE_GOYONG) * ratio)
                if need < SAFETY_MIN_DIRECT:
                    need = (target - tax) / (1 + 0.08 + (RATE_SANJAE + RATE_GOYONG) * ratio)
                k = need / direct
                _scale(sections, k)
                scale_k *= k
                direct_mat, direct_lab = _sum_items(sections)
                direct = direct_mat + direct_lab
                sj, gy, sf = _lotte_indirect(direct, direct_lab)
                misc = target - (direct + sj + gy + sf) - tax
                if 0 <= misc <= direct * RATE_MISC_MAX:
                    break
        else:
            k = target / direct
            _scale(sections, k)
            scale_k *= k
            direct_mat, direct_lab = _sum_items(sections)
            direct = direct_mat + direct_lab
            resid = target - direct
            it = _fit_item(sections, proj.get("fit_item", "현장정리"))
            qty = float(it.get("qty", 1))
            add = int(round(resid / qty))
            it["lab"] = int(it.get("lab", 0) or 0) + add
            if it["lab"] < 0:
                it["mat"] = int(it.get("mat", 0) or 0) + it["lab"]; it["lab"] = 0
            direct_mat, direct_lab = _sum_items(sections)
            direct = direct_mat + direct_lab
            # 수량 반올림 잔차가 남으면 한 번 더
            resid = target - direct
            if resid:
                it["lab"] += resid if qty == 1 else 0
                direct_mat, direct_lab = _sum_items(sections)
                direct = direct_mat + direct_lab
        fit_note = f"단가 {scale_k:.3f}배 비례조정 후 목표에 맞춤"
        if scale_k < 0.75 or scale_k > 1.35:
            fit_note += " (조정폭이 큼 → 항목 구성/수량 재검토 권장)"

    res = {"mode": mode, "direct_mat": direct_mat, "direct_lab": direct_lab, "direct": direct, "scale": round(scale_k, 3),
           "indirect": [], "fit": {"target": target, "ok": True, "message": ""}}

    if mode == "lotte":
        sanjae, goyong, safety = _lotte_indirect(direct, direct_lab)
        base = direct + sanjae + goyong + safety
        if target:
            tax = stamp_tax(target)
            misc = target - base - tax
            misc_max = int(direct * RATE_MISC_MAX)
            if misc < 0 or misc > misc_max:
                res["fit"]["ok"] = False
                need_lo = target - tax - (base - direct) - misc_max
                need_hi = target - tax - (base - direct)
                res["fit"]["message"] = (f"목표 {target:,}원에 맞추려면 공과잡비가 {misc:,}원이어야 하는데 허용범위(0 ~ {misc_max:,})를 벗어남. "
                                         f"직접공사비를 {need_lo:,} ~ {need_hi:,}원 사이로 조정 필요 (현재 {direct:,}).")
                misc = min(max(misc, 0), misc_max)
        else:
            misc = int(direct * RATE_MISC_MAX)
            tax = stamp_tax((base + misc) // 10000 * 10000)
        indirect_sum = sanjae + goyong + safety + misc + tax
        total = (direct + indirect_sum) // 10000 * 10000
        res["indirect"] = [
            ("산재보험료", f"노무비의 {RATE_SANJAE*100:.2f}%", sanjae, "전체공사 적용"),
            ("고용보험료", f"노무비의 {RATE_GOYONG*100:.2f}%", goyong, ""),
            ("안전관리비", f"직접공사비의 {RATE_SAFETY*100:.2f}%", safety, "2천만원 이상 공사" if safety else "2천만원 미만 미적용"),
            ("공과잡비", f"직접공사비의 {RATE_MISC_MAX*100:.0f}% 이내", misc, f"기타항목 포함 ({misc/direct*100:.2f}%)" if direct else ""),
            ("인지세", "당사 분담분 (VAT 별도)", tax, "양사 5:5 분담"),
        ]
        res["indirect_sum"] = indirect_sum
    else:
        if target and not (target <= direct < target + 10000):
            res["fit"]["ok"] = False
            res["fit"]["message"] = (f"목표 {target:,}원. 직접공사비를 {target:,} ~ {target+9999:,}원 사이로 맞춰야 함 "
                                     f"(현재 {direct:,}, 차이 {direct-target:+,}).")
        total = direct // 10000 * 10000
        res["indirect_sum"] = 0

    vat = int(round(total * 0.1))
    res.update({"total": total, "vat": vat, "grand": total + vat, "sections": sections, "total_kor": korean_money(total)})
    if target and res["fit"]["ok"]:
        res["fit"]["message"] = f"목표 {target:,}원 일치" + (f" · {fit_note}" if fit_note else "")
    elif fit_note:
        res["fit"]["message"] += " · " + fit_note
    return res


# ───────────────────────────── PDF ─────────────────────────────
class SignBlock(Flowable):
    """발주자 / 시공자 서명란 (도장 자동 날인)"""
    def __init__(self, width, client, stamp_path, fonts):
        super().__init__()
        self.width, self.client, self.stamp, self.fonts = width, client, stamp_path, fonts
        self.height = 30 * mm

    def wrap(self, aw, ah):
        return self.width, self.height

    def draw(self):
        c = self.canv
        R, B = self.fonts
        w, h = self.width, self.height
        c.setStrokeColor(LINE_DARK); c.setLineWidth(0.8)
        c.rect(0, 0, w, h)
        c.line(w / 2, 0, w / 2, h)
        c.setFillColor(GRAY_FILL2)
        c.rect(0, h - 7 * mm, w, 7 * mm, stroke=1, fill=1)
        c.setFillColor(DARK); c.setFont(B, 10)
        c.drawCentredString(w / 4, h - 5 * mm, "발  주  자")
        c.drawCentredString(3 * w / 4, h - 5 * mm, "시  공  자")
        # 발주자 / 시공자 — 라벨 x=8mm, 값 x=24mm 로 통일 (성명 밑줄도 24mm 에서 시작)
        c.setFont(R, 9.5)
        LX, VX = 8 * mm, 24 * mm
        c.drawString(LX, h - 14 * mm, "상  호 :"); c.drawString(VX, h - 14 * mm, self.client)
        c.drawString(LX, h - 22.5 * mm, "성  명 :")
        c.line(VX, h - 23.5 * mm, w / 2 - 20 * mm, h - 23.5 * mm)
        c.drawString(w / 2 - 17 * mm, h - 22.5 * mm, "(인)")
        x0 = w / 2
        c.drawString(x0 + LX, h - 14 * mm, "상  호 :"); c.drawString(x0 + VX, h - 14 * mm, COMPANY['name'])
        c.drawString(x0 + LX, h - 22.5 * mm, "성  명 :"); c.drawString(x0 + VX, h - 22.5 * mm, COMPANY['ceo'])
        c.drawString(w - 17 * mm, h - 22.5 * mm, "(인)")
        if self.stamp:
            s = 19 * mm
            c.drawImage(ImageReader(self.stamp), w - 23 * mm, h - 28.5 * mm, s, s, mask="auto")


class SubmitBlock(Flowable):
    """마지막 장: 제출 문구 + 도장"""
    def __init__(self, width, date_str, stamp_path, fonts):
        super().__init__()
        self.width, self.date_str, self.stamp, self.fonts = width, date_str, stamp_path, fonts
        self.height = 40 * mm

    def wrap(self, aw, ah):
        return self.width, self.height

    def draw(self):
        c = self.canv
        R, B = self.fonts
        w, h = self.width, self.height
        c.setFillColor(DARK); c.setFont(B, 12)
        c.drawCentredString(w / 2, h - 10 * mm, "상기와 같이 견적서를 제출합니다.")
        c.setFont(R, 10.5)
        c.drawCentredString(w / 2, h - 19 * mm, self.date_str)
        c.setFont(B, 12)
        txt = f"{COMPANY['name']}    대표   {COMPANY['ceo']}    (인)"
        c.drawCentredString(w / 2, h - 32 * mm, txt)
        if self.stamp:
            tw = c.stringWidth(txt, B, 12)
            s = 20 * mm
            c.drawImage(ImageReader(self.stamp), w / 2 + tw / 2 - 11 * mm, h - 41 * mm, s, s, mask="auto")


def transparent_asset(path):
    """도장·로고 흰 배경 강제 제거 — 파일이 어디서 왔든(프로젝트·업로드·깃허브) 항상 투명 PNG 로 만들어 쓴다.
    이미 투명하면 그대로 쓰고, 흰 배경이면 color-to-alpha 방식으로 흰색만 빼서 ./assets/_t_<이름> 에 저장한다."""
    if not path or not os.path.isfile(path):
        return path
    try:
        im = PILImage.open(path).convert("RGBA")
        w, h = im.size
        px = im.load()
        corners = [px[0, 0], px[w - 1, 0], px[0, h - 1], px[w - 1, h - 1]]
        if all(c[3] < 20 for c in corners):
            return path  # 이미 배경이 투명
        import numpy as np
        arr = np.asarray(im).astype(np.float64)
        rgb, a0 = arr[..., :3], arr[..., 3]
        al = (255.0 - rgb.min(axis=2)) / 255.0
        safe = np.where(al > 0, al, 1.0)[..., None]
        col = np.clip((rgb - 255.0 * (1.0 - al[..., None])) / safe, 0, 255)
        alpha = np.where(al < 0.06, 0, np.round(al * a0))
        col[alpha == 0] = 255
        out = PILImage.fromarray(np.dstack([col, alpha]).astype(np.uint8), "RGBA")
        os.makedirs("./assets", exist_ok=True)
        dst = os.path.abspath(os.path.join("./assets", "_t_" + os.path.basename(path)))
        out.save(dst)
        return dst
    except Exception:
        return path


def make_watermark(logo_path, alpha=0.06):
    if not logo_path:
        return None
    im = PILImage.open(logo_path).convert("RGBA")
    r, g, b, a = im.split()
    a = a.point(lambda v: int(v * alpha))
    im.putalpha(a)
    os.makedirs("./assets", exist_ok=True)
    out = os.path.abspath(os.path.join("./assets", "_wm.png"))
    im.save(out)
    return out


def build_pdf(data, res, out_path, assets_dir):
    reg, bold, idx = find_font(assets_dir)
    pdfmetrics.registerFont(TTFont("KR", reg, subfontIndex=idx))
    pdfmetrics.registerFont(TTFont("KRB", bold, subfontIndex=idx))
    FONTS = ("KR", "KRB")
    logo = transparent_asset(find_asset(LOGO_FILE, assets_dir))
    stamp = transparent_asset(find_asset(STAMP_FILE, assets_dir))
    wm = make_watermark(logo)
    proj = data["project"]
    mode = res["mode"]
    date_str = proj.get("date") or datetime.date.today().isoformat()
    y, m, d = date_str.split("-")
    date_kor = f"{y}년 {int(m)}월 {int(d)}일"

    PW, PH = A4
    M = 13 * mm

    def on_page(canvas, doc):
        w, h = PW, PH
        canvas.saveState()
        if wm:
            ww = 120 * mm
            im = PILImage.open(wm)
            hh = ww * im.size[1] / im.size[0]
            canvas.drawImage(ImageReader(wm), (w - ww) / 2, (h - hh) / 2, ww, hh, mask="auto")
        # 상단 얇은 골드 라인
        canvas.setStrokeColor(GOLD); canvas.setLineWidth(2.2)
        canvas.line(0, h - 6 * mm, w, h - 6 * mm)
        # 푸터
        canvas.setFont("KR", 7.5); canvas.setFillColor(GRAY_TXT)
        canvas.drawString(M, 8 * mm, f"{COMPANY['name']}  |  {COMPANY['address']}  |  TEL {COMPANY['tel']}")
        canvas.drawRightString(w - M, 8 * mm, f"-  {doc.page}  -")
        if doc.page > 1 and logo:
            im = PILImage.open(logo)
            lw = 14 * mm
            canvas.drawImage(ImageReader(logo), w - M - lw, h - 6 * mm - 9 * mm, lw, lw * im.size[1] / im.size[0], mask="auto")
        canvas.restoreState()

    doc = BaseDocTemplate(out_path, pagesize=A4, leftMargin=M, rightMargin=M, topMargin=14 * mm, bottomMargin=16 * mm,
                          title=f"견적서 - {proj.get('site_name','')}", author=COMPANY["name"])
    fr_p = Frame(M, 16 * mm, PW - 2 * M, PH - 30 * mm, id="p", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    fr_s = Frame(M, 16 * mm, PW - 2 * M, PH - 34 * mm, id="s", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates([
        PageTemplate(id="cover", frames=[fr_p], pagesize=A4, onPage=on_page),
        PageTemplate(id="detail", frames=[fr_s], pagesize=A4, onPage=on_page),
    ])

    W = PW - 2 * M
    def st(name, size, **kw):
        kw.setdefault("leading", size * 1.35)
        return ParagraphStyle(name, fontName=kw.pop("font", "KR"), fontSize=size, **kw)
    S_TITLE = st("t", 22, font="KRB", alignment=TA_CENTER, textColor=DARK, leading=27)
    S_H = st("h", 11, font="KRB", textColor=NAVY)
    S_SM = st("sm", 8.5, textColor=DARK)
    S_SMG = st("smg", 8, textColor=GRAY_TXT)
    S_CELL = st("cell", 7, textColor=DARK, leading=9)
    S_CELLB = st("cellb", 7.5, font="KRB", textColor=DARK)
    S_CELL_C = st("cellc", 8, textColor=DARK, alignment=TA_CENTER)
    S_NOTE = st("note", 6.8, textColor=colors.HexColor("#B03A00"), leading=8.5)
    S_COND = st("cond", 9.5, textColor=DARK, leading=15)

    story = []
    # ── 표지 헤더
    if logo:
        im = PILImage.open(logo)
        lw = 34 * mm
        from reportlab.platypus import Image as RLImage
        logo_fl = RLImage(logo, lw, lw * im.size[1] / im.size[0])
    else:
        logo_fl = Paragraph("UNION ONE", st("lgt", 16, font="KRB", textColor=NAVY))
    hdr = Table([[logo_fl, ""]], colWidths=[W * 0.5, W * 0.5])
    hdr.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story += [hdr, Spacer(1, 2.5 * mm)]
    story.append(Paragraph("견 &nbsp; 적 &nbsp; 서", S_TITLE))
    rule = Table([[""]], colWidths=[W], rowHeights=[1.2 * mm])
    rule.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), GOLD)]))
    story += [Spacer(1, 1 * mm), rule, Spacer(1, 3.5 * mm)]

    # ── 공급자 / 공사개요
    def info_table(title, rows, width):
        body = [[Paragraph(f"<b>{title}</b>", st("it", 9, font="KRB", textColor=colors.white)), ""]]
        for k, v in rows:
            body.append([Paragraph(k, S_CELL_C), Paragraph(str(v), S_CELL)])
        t = Table(body, colWidths=[22 * mm, width - 22 * mm], rowHeights=[5.8 * mm] + [None] * len(rows))
        t.setStyle(TableStyle([
            ("SPAN", (0, 0), (1, 0)), ("BACKGROUND", (0, 0), (1, 0), NAVY),
            ("BACKGROUND", (0, 1), (0, -1), GRAY_FILL),
            ("GRID", (0, 0), (-1, -1), 0.5, LINE), ("BOX", (0, 0), (-1, -1), 0.8, LINE_DARK),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 1), (-1, -1), 1.6), ("BOTTOMPADDING", (0, 1), (-1, -1), 1.6),
        ]))
        return t

    left = info_table("공 급 자", [
        ("등록번호", COMPANY["reg_no"]), ("상 호", COMPANY["name"]), ("대표자", COMPANY["ceo"]),
        ("소재지", COMPANY["address"]), ("업태/종목", f"{COMPANY['biz_type']} / {COMPANY['biz_item']}"),
        ("전 화", COMPANY["tel"]),
    ], W * 0.5 - 2 * mm)
    right_rows = [
        ("수신처", proj.get("client", "")), ("현장명", proj.get("site_name", "")),
        ("현장주소", proj.get("site_address", "")), ("층 / 면적", f"{proj.get('floor','')}  /  {proj.get('area_py','')}평".strip(" /")),
        ("견적일자", date_str), ("공사기간", f"착공일로부터 {proj.get('duration_days','')}일"),
        ("견적유효", f"견적일로부터 {proj.get('valid_days',30)}일"),
    ]
    rightT = info_table("공 사 개 요", right_rows, W * 0.5 - 2 * mm)
    two = Table([[left, rightT]], colWidths=[W * 0.5, W * 0.5])
    two.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story += [two, Spacer(1, 3.5 * mm)]

    # ── 견적금액
    amt = Table([[Paragraph("견 적 금 액", st("a1", 10, font="KRB", textColor=colors.white, alignment=TA_CENTER)),
                  Paragraph(f"일금 &nbsp;<b>{res['total_kor']}원정</b> &nbsp;( ₩ {res['total']:,} ) &nbsp;<font size=8 color='#666666'>부가세 별도</font>",
                            st("a2", 12.5, textColor=DARK)),
                  Paragraph(f"<font size=7.5 color='#666666'>부가세 포함</font><br/><b>₩ {res['grand']:,}</b>", st("a3", 10, alignment=TA_RIGHT, textColor=DARK, leading=13))]],
                colWidths=[28 * mm, W - 28 * mm - 42 * mm, 42 * mm], rowHeights=[11 * mm])
    amt.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, 0), NAVY), ("BACKGROUND", (1, 0), (-1, 0), NAVY_LIGHT),
                             ("BOX", (0, 0), (-1, -1), 1, NAVY), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                             ("LEFTPADDING", (1, 0), (1, 0), 6), ("RIGHTPADDING", (2, 0), (2, 0), 6)]))
    story += [amt, Spacer(1, 3.5 * mm)]

    # ── 총괄 집계
    story.append(Paragraph("공 사 비 총 괄", S_H))
    story.append(Spacer(1, 1.5 * mm))
    cw = [W * 0.31, W * 0.16, W * 0.16, W * 0.17, W * 0.20]
    hdr_row = ["구  분", "재 료 비", "노 무 비", "합  계", "비  고"]
    rows = [hdr_row]
    style = [("BACKGROUND", (0, 0), (-1, 0), GRAY_FILL2), ("FONTNAME", (0, 0), (-1, 0), "KRB"),
             ("FONTNAME", (0, 1), (-1, -1), "KR"), ("FONTSIZE", (0, 0), (-1, -1), 8),
             ("ALIGN", (1, 0), (3, -1), "RIGHT"), ("ALIGN", (0, 0), (-1, 0), "CENTER"), ("ALIGN", (4, 0), (4, -1), "CENTER"),
             ("GRID", (0, 0), (-1, -1), 0.4, LINE), ("BOX", (0, 0), (-1, -1), 0.8, LINE_DARK),
             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
             ("TEXTCOLOR", (0, 0), (-1, -1), DARK)]
    for sec in res["sections"]:
        rows.append([f"  {sec['_no']}. {sec['name']}", won(sec["_mat"]), won(sec["_lab"]), won(sec["_sum"]), Paragraph(sec.get("note", "") or "", S_CELL_C)])
    r = len(rows)
    rows.append(["직접공사비 계", won(res["direct_mat"]), won(res["direct_lab"]), won(res["direct"]), ""])
    style += [("BACKGROUND", (0, r), (-1, r), GRAY_FILL), ("FONTNAME", (0, r), (-1, r), "KRB"), ("ALIGN", (0, r), (0, r), "CENTER")]
    if mode == "lotte":
        for name, spec, amtv, note in res["indirect"]:
            rows.append([f"  {name}  ({spec})", "", "", won(amtv), Paragraph(note, S_CELL_C)])
        r = len(rows)
        rows.append(["간접공사비 계", "", "", won(res["indirect_sum"]), ""])
        style += [("BACKGROUND", (0, r), (-1, r), GRAY_FILL), ("FONTNAME", (0, r), (-1, r), "KRB"), ("ALIGN", (0, r), (0, r), "CENTER")]
    r = len(rows)
    rows.append(["합 계 (부가세 별도)", "", "", won(res["total"]), Paragraph("만 단위 절사", st("wn", 8, textColor=colors.white, alignment=TA_CENTER))])
    style += [("FONTNAME", (0, r), (-1, r), "KRB"), ("ALIGN", (0, r), (0, r), "CENTER"),
              ("BACKGROUND", (0, r), (-1, r), NAVY), ("TEXTCOLOR", (0, r), (-1, r), colors.white), ("FONTSIZE", (0, r), (-1, r), 9.5),
              ("TOPPADDING", (0, r), (-1, r), 2.5), ("BOTTOMPADDING", (0, r), (-1, r), 2.5)]
    t = Table(rows, colWidths=cw)
    t.setStyle(TableStyle(style))
    story += [t, Spacer(1, 3 * mm)]
    story.append(Paragraph("위 견적 내용을 확인하였으며, 발주자와 시공자는 상기 금액 및 부대조건(별지)에 합의합니다.  <font color='#666666' size=7.5>※ 세부 내역은 「공사 내역서」 참조</font>", S_SM))
    story.append(Spacer(1, 1.5 * mm))
    story.append(SignBlock(W, proj.get("client", ""), stamp, FONTS))

    # ── 내역서 (가로)
    story.append(NextPageTemplate("detail"))
    story.append(PageBreak())
    story.append(Paragraph("공 사 내 역 서", st("dt", 14, font="KRB", textColor=DARK)))
    story.append(Paragraph(f"공사명 : {proj.get('site_name','')}", S_SMG))
    story.append(Spacer(1, 2.5 * mm))
    cws = [31, 35, 8, 10.5, 12.5, 15.5, 12.5, 15.5, 12.5, 16.5, 14.5]
    scale = W / (sum(cws) * mm)
    cws = [c * mm * scale for c in cws]
    head1 = ["품  명", "규  격", "단위", "수량", "재 료 비", "", "노 무 비", "", "합  계", "", "비  고"]
    head2 = ["", "", "", "", "단 가", "금 액", "단 가", "금 액", "단 가", "금 액", ""]
    drows = [head1, head2]
    dstyle = [("SPAN", (0, 0), (0, 1)), ("SPAN", (1, 0), (1, 1)), ("SPAN", (2, 0), (2, 1)), ("SPAN", (3, 0), (3, 1)),
              ("SPAN", (4, 0), (5, 0)), ("SPAN", (6, 0), (7, 0)), ("SPAN", (8, 0), (9, 0)), ("SPAN", (10, 0), (10, 1)),
              ("BACKGROUND", (0, 0), (-1, 1), GRAY_FILL2), ("FONTNAME", (0, 0), (-1, 1), "KRB"), ("ALIGN", (0, 0), (-1, 1), "CENTER"),
              ("FONTNAME", (0, 2), (-1, -1), "KR"), ("FONTSIZE", (0, 0), (-1, -1), 7), ("TEXTCOLOR", (0, 0), (-1, -1), DARK),
              ("ALIGN", (2, 2), (3, -1), "CENTER"), ("ALIGN", (4, 2), (9, -1), "RIGHT"),
              ("GRID", (0, 0), (-1, -1), 0.35, LINE), ("BOX", (0, 0), (-1, -1), 0.8, LINE_DARK), ("LINEBELOW", (0, 1), (-1, 1), 0.8, LINE_DARK),
              ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
              ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2)]
    for sec in res["sections"]:
        i = len(drows)
        drows.append([Paragraph(f"<b>{sec['_no']}. {sec['name']}</b>", S_CELLB)] + [""] * 10)
        dstyle += [("SPAN", (0, i), (-1, i)), ("BACKGROUND", (0, i), (-1, i), NAVY_LIGHT), ("NOSPLIT", (0, i), (-1, i + 1))]
        for it in sec["items"]:
            drows.append([Paragraph(it.get("name", ""), S_CELL), Paragraph(it.get("spec", "") or "", S_CELL), it.get("unit", ""), fmt_qty(it["_qty"]),
                          won(it["_mat"]), won(it["_mat_amt"]), won(it["_lab"]), won(it["_lab_amt"]), won(it["_sum"]), won(it["_sum_amt"]),
                          Paragraph(it.get("note", "") or "", S_NOTE if it.get("note") else S_CELL)])
        i = len(drows)
        drows.append(["", "소   계", "", "", "", won(sec["_mat"]), "", won(sec["_lab"]), "", won(sec["_sum"]), ""])
        dstyle += [("BACKGROUND", (0, i), (-1, i), GRAY_FILL), ("FONTNAME", (0, i), (-1, i), "KRB"), ("ALIGN", (1, i), (1, i), "CENTER"),
                   ("NOSPLIT", (0, i - 1), (-1, i))]
    i = len(drows)
    drows.append(["", "직 접 공 사 비  계", "", "", "", won(res["direct_mat"]), "", won(res["direct_lab"]), "", won(res["direct"]), ""])
    dstyle += [("BACKGROUND", (0, i), (-1, i), GRAY_FILL2), ("FONTNAME", (0, i), (-1, i), "KRB"), ("ALIGN", (1, i), (1, i), "CENTER"),
               ("LINEABOVE", (0, i), (-1, i), 0.8, LINE_DARK), ("NOSPLIT", (0, i - 1), (-1, i))]
    dt = Table(drows, colWidths=cws, repeatRows=2)
    dt.setStyle(TableStyle(dstyle))
    story.append(dt)

    # ── 부대조건 (세로)
    story.append(PageBreak())
    story.append(Paragraph("부대조건 및 특기사항", st("tt", 14, font="KRB", textColor=DARK)))
    story += [Spacer(1, 1.5 * mm), rule, Spacer(1, 5 * mm)]
    conds = data.get("conditions") or DEFAULT_CONDITIONS
    crow = []
    for n, c in enumerate(conds, 1):
        crow.append([Paragraph(f"<b>{n}.</b>", st("cn", 9.5, font="KRB", textColor=GOLD)), Paragraph(c, S_COND)])
    ct = Table(crow, colWidths=[8 * mm, W - 8 * mm])
    ct.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                            ("LINEBELOW", (0, 0), (-1, -2), 0.3, LINE)]))
    story += [ct, Spacer(1, 8 * mm)]
    bank = Table([[Paragraph("입 금 계 좌", st("b1", 9.5, font="KRB", textColor=colors.white, alignment=TA_CENTER)),
                   Paragraph(f"<b>{COMPANY['bank']}</b>", st("b2", 10.5, textColor=DARK))]], colWidths=[30 * mm, W - 30 * mm], rowHeights=[10 * mm])
    bank.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, 0), NAVY), ("BACKGROUND", (1, 0), (1, 0), NAVY_LIGHT), ("BOX", (0, 0), (-1, -1), 1, NAVY),
                              ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (1, 0), (1, 0), 8)]))
    story += [bank, Spacer(1, 16 * mm), SubmitBlock(W, date_kor, stamp, FONTS)]
    doc.build(story)


# ───────────────────────────── XLSX ─────────────────────────────
def build_xlsx(data, res, out_path, assets_dir):
    proj = data["project"]
    mode = res["mode"]
    date_str = proj.get("date") or datetime.date.today().isoformat()
    y, m, d = date_str.split("-")
    wb = openpyxl.Workbook()
    F = "맑은 고딕"
    thin = Side(style="thin", color="BFBFBF"); med = Side(style="medium", color="555555")
    B_ALL = Border(left=thin, right=thin, top=thin, bottom=thin)
    OR = PatternFill("solid", fgColor="1B2A41"); ORL = PatternFill("solid", fgColor="EEF1F5"); GD = PatternFill("solid", fgColor="B08D57")
    G1 = PatternFill("solid", fgColor="F3F3F3"); G2 = PatternFill("solid", fgColor="E6E6E6")
    NUM = '#,##0;-#,##0;"-"'
    C = Alignment(horizontal="center", vertical="center", wrap_text=True)
    L = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    Rt = Alignment(horizontal="right", vertical="center")
    f = lambda **kw: Font(name=F, size=kw.pop("size", 9), **kw)

    # ── 내역서 sheet 먼저 (표지가 참조)
    ws2 = wb.active; ws2.title = "내역서"
    widths = [24, 34, 6, 8, 11, 13, 11, 13, 11, 14, 22]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.merge_cells("A1:K1"); ws2["A1"] = "공  사  내  역  서"; ws2["A1"].font = f(size=16, bold=True); ws2["A1"].alignment = C
    ws2.row_dimensions[1].height = 30
    ws2["A2"] = f"공사명 : {proj.get('site_name','')}"; ws2["A2"].font = f(size=9)
    ws2.merge_cells("A2:F2")
    ws2["G2"] = f"수신처 : {proj.get('client','')}"; ws2["G2"].font = f(size=9); ws2.merge_cells("G2:K2"); ws2["G2"].alignment = Rt
    hdr1 = ["품  명", "규  격", "단위", "수량", "재 료 비", "", "노 무 비", "", "합  계", "", "비  고"]
    hdr2 = ["", "", "", "", "단 가", "금 액", "단 가", "금 액", "단 가", "금 액", ""]
    for j, (a, b) in enumerate(zip(hdr1, hdr2), 1):
        ws2.cell(4, j, a); ws2.cell(5, j, b)
    for rng in ["A4:A5", "B4:B5", "C4:C5", "D4:D5", "E4:F4", "G4:H4", "I4:J4", "K4:K5"]:
        ws2.merge_cells(rng)
    for r in (4, 5):
        for j in range(1, 12):
            c = ws2.cell(r, j); c.font = f(bold=True); c.fill = G2; c.alignment = C; c.border = B_ALL
    ws2.row_dimensions[4].height = 18; ws2.row_dimensions[5].height = 18
    row = 6
    sub_rows = []
    for sec in res["sections"]:
        ws2.cell(row, 1, f"{sec['_no']}. {sec['name']}").font = f(bold=True)
        ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
        for j in range(1, 12):
            ws2.cell(row, j).fill = ORL; ws2.cell(row, j).border = B_ALL
        ws2.cell(row, 1).alignment = L
        row += 1
        first = row
        for it in sec["items"]:
            vals = [it.get("name", ""), it.get("spec", "") or "", it.get("unit", ""), it["_qty"], it["_mat"], f"=E{row}*D{row}",
                    it["_lab"], f"=G{row}*D{row}", f"=E{row}+G{row}", f"=F{row}+H{row}", it.get("note", "") or ""]
            for j, v in enumerate(vals, 1):
                c = ws2.cell(row, j, v); c.font = f(); c.border = B_ALL
                c.alignment = L if j in (1, 2, 11) else (C if j in (3,) else Rt)
                if j >= 4:
                    c.number_format = NUM if j != 4 else "#,##0.##"
                if j in (4, 5, 7):
                    c.font = f(color="0000FF")
                if j == 11 and v:
                    c.font = f(color="B03A00")
            row += 1
        last = row - 1
        ws2.cell(row, 2, "소     계").font = f(bold=True)
        for col in ("F", "H", "J"):
            ws2[f"{col}{row}"] = f"=SUM({col}{first}:{col}{last})" if last >= first else 0
            ws2[f"{col}{row}"].number_format = NUM; ws2[f"{col}{row}"].font = f(bold=True); ws2[f"{col}{row}"].alignment = Rt
        for j in range(1, 12):
            ws2.cell(row, j).fill = G1; ws2.cell(row, j).border = B_ALL
        ws2.cell(row, 2).alignment = C
        sec["_xl_row"] = row
        sub_rows.append(row)
        row += 1
    ws2.cell(row, 2, "직 접 공 사 비  계").font = f(bold=True); ws2.cell(row, 2).alignment = C
    for col in ("F", "H", "J"):
        ws2[f"{col}{row}"] = "=" + "+".join(f"{col}{r}" for r in sub_rows)
        ws2[f"{col}{row}"].number_format = NUM; ws2[f"{col}{row}"].font = f(bold=True); ws2[f"{col}{row}"].alignment = Rt
    for j in range(1, 12):
        ws2.cell(row, j).fill = G2; ws2.cell(row, j).border = B_ALL
    direct_row = row
    ws2.freeze_panes = "A6"
    ws2.print_title_rows = "4:5"
    ws2.page_setup.orientation = "landscape"; ws2.page_setup.fitToWidth = 1; ws2.page_setup.fitToHeight = 0
    ws2.sheet_properties.pageSetUpPr.fitToPage = True
    ws2.print_options.horizontalCentered = True
    ws2.page_margins.left = ws2.page_margins.right = 0.4

    # ── 견적서(표지) sheet
    ws = wb.create_sheet("견적서", 0)
    for col, w in zip("ABCDEFGH", [13, 22, 14, 3, 13, 22, 14, 3]):
        ws.column_dimensions[col].width = w
    logo = transparent_asset(find_asset(LOGO_FILE, assets_dir))
    if logo:
        img = XLImage(logo); ratio = img.height / img.width; img.width = 120; img.height = int(120 * ratio)
        ws.add_image(img, "A1")
    ws.row_dimensions[1].height = 48
    ws.merge_cells("A3:H3"); ws["A3"] = "견   적   서"; ws["A3"].font = f(size=22, bold=True); ws["A3"].alignment = C
    ws.row_dimensions[3].height = 40
    for j in range(1, 9):
        ws.cell(4, j).fill = GD
    ws.row_dimensions[4].height = 4

    def block(r0, c0, title, rows):
        ws.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c0 + 2)
        c = ws.cell(r0, c0, title); c.font = f(bold=True, color="FFFFFF"); c.fill = OR; c.alignment = C
        for j in range(3):
            ws.cell(r0, c0 + j).border = B_ALL
        for k, (a, b) in enumerate(rows, 1):
            ws.cell(r0 + k, c0, a).font = f(); ws.cell(r0 + k, c0).fill = G1; ws.cell(r0 + k, c0).alignment = C
            ws.merge_cells(start_row=r0 + k, start_column=c0 + 1, end_row=r0 + k, end_column=c0 + 2)
            ws.cell(r0 + k, c0 + 1, b).font = f(); ws.cell(r0 + k, c0 + 1).alignment = L
            for j in range(3):
                ws.cell(r0 + k, c0 + j).border = B_ALL
            ws.row_dimensions[r0 + k].height = 18
    block(6, 1, "공  급  자", [("등록번호", COMPANY["reg_no"]), ("상 호", COMPANY["name"]), ("대표자", COMPANY["ceo"]),
                            ("소재지", COMPANY["address"]), ("업태/종목", f"{COMPANY['biz_type']} / {COMPANY['biz_item']}"),
                            ("전 화", COMPANY["tel"])])
    block(6, 5, "공  사  개  요", [("수신처", proj.get("client", "")), ("현장명", proj.get("site_name", "")), ("현장주소", proj.get("site_address", "")),
                              ("층 / 면적", f"{proj.get('floor','')} / {proj.get('area_py','')}평"), ("견적일자", date_str),
                              ("공사기간", f"착공일로부터 {proj.get('duration_days','')}일"), ("견적유효", f"견적일로부터 {proj.get('valid_days',30)}일")])
    ws.row_dimensions[6].height = 20; ws.row_dimensions[9].height = 30
    # 금액
    r = 15
    ws.merge_cells(f"A{r}:A{r}"); ws[f"A{r}"] = "견적금액"; ws[f"A{r}"].font = f(bold=True, color="FFFFFF"); ws[f"A{r}"].fill = OR; ws[f"A{r}"].alignment = C
    ws.merge_cells(f"B{r}:F{r}")
    ws[f"B{r}"].font = f(size=12, bold=True); ws[f"B{r}"].fill = ORL; ws[f"B{r}"].alignment = L
    ws.merge_cells(f"G{r}:H{r}"); ws[f"G{r}"].fill = ORL; ws[f"G{r}"].font = f(size=11, bold=True); ws[f"G{r}"].alignment = Rt; ws[f"G{r}"].number_format = '"₩"#,##0'
    ws.row_dimensions[r].height = 34
    for j in range(1, 9):
        ws.cell(r, j).border = Border(top=med, bottom=med, left=med if j == 1 else None, right=med if j == 8 else None)
    # 총괄
    r = 18
    ws[f"A{r}"] = "공 사 비 총 괄"; ws[f"A{r}"].font = f(size=10, bold=True, color="1B2A41")
    r += 1
    heads = ["구  분", "", "재 료 비", "", "노 무 비", "합  계", "비  고", ""]
    for j, h in enumerate(heads, 1):
        c = ws.cell(r, j, h); c.font = f(bold=True); c.fill = G2; c.alignment = C; c.border = B_ALL
    ws.merge_cells(f"A{r}:B{r}"); ws.merge_cells(f"C{r}:D{r}"); ws.merge_cells(f"G{r}:H{r}")
    r += 1

    def trow(r, label, mat, lab, tot, note, fill=None, bold=False, white=False):
        ws.merge_cells(f"A{r}:B{r}"); ws.merge_cells(f"C{r}:D{r}"); ws.merge_cells(f"G{r}:H{r}")
        vals = {1: label, 3: mat, 5: lab, 6: tot, 7: note}
        for j in range(1, 9):
            c = ws.cell(r, j); c.border = B_ALL
            if j in vals:
                c.value = vals[j]
            c.font = f(bold=bold, color="FFFFFF" if white else "000000")
            c.alignment = L if j == 1 else (C if j == 7 else Rt)
            if j in (3, 5, 6):
                c.number_format = NUM
            if fill:
                c.fill = fill
        ws.row_dimensions[r].height = 17

    for sec in res["sections"]:
        xr = sec["_xl_row"]
        trow(r, f"  {sec['_no']}. {sec['name']}", f"=내역서!F{xr}", f"=내역서!H{xr}", f"=내역서!J{xr}", sec.get("note", ""))
        r += 1
    trow(r, "직접공사비 계", f"=내역서!F{direct_row}", f"=내역서!H{direct_row}", f"=내역서!J{direct_row}", "", G1, True)
    DR = r; r += 1
    if mode == "lotte":
        ind = res["indirect"]
        trow(r, f"  산재보험료 (노무비의 {RATE_SANJAE*100:.2f}%)", None, None, f"=ROUND(E{DR}*{RATE_SANJAE},0)", ind[0][3]); r += 1
        trow(r, f"  고용보험료 (노무비의 {RATE_GOYONG*100:.2f}%)", None, None, f"=ROUND(E{DR}*{RATE_GOYONG},0)", ind[1][3]); r += 1
        trow(r, f"  안전관리비 (직접공사비의 {RATE_SAFETY*100:.2f}%)", None, None, f"=IF(F{DR}>={SAFETY_MIN_DIRECT},ROUND(F{DR}*{RATE_SAFETY},0),0)", ind[2][3]); r += 1
        trow(r, "  공과잡비 (직접공사비의 10% 이내)", None, None, ind[3][2], ind[3][3]); ws[f"F{r}"].font = f(color="0000FF"); r += 1
        trow(r, "  인지세 (당사 분담분)", None, None, ind[4][2], ind[4][3]); ws[f"F{r}"].font = f(color="0000FF"); r += 1
        trow(r, "간접공사비 계", None, None, f"=SUM(F{DR+1}:F{r-1})", "", G1, True); IR = r; r += 1
        trow(r, "합 계 (부가세 별도)", None, None, f"=ROUNDDOWN(F{DR}+F{IR},-4)", "만 단위 절사", OR, True, True)
    else:
        trow(r, "합 계 (부가세 별도)", None, None, f"=ROUNDDOWN(F{DR},-4)", "만 단위 절사", OR, True, True)
    TR = r; r += 1
    ws[f"B15"] = f"일금 {res['total_kor']}원정 (부가세 별도)"
    ws[f"G15"] = f"=F{TR}"
    ws.merge_cells("G16:H16"); ws["G16"] = f'="부가세 포함  ₩"&TEXT(F{TR}*1.1,"#,##0")'; ws["G16"].font = f(size=8, color="666666"); ws["G16"].alignment = Rt
    ws[f"A{r}"] = "※ 파란색 숫자는 입력값, 검은색은 수식입니다. 한글 금액 표기는 내역 변경 시 직접 수정해 주세요."; ws[f"A{r}"].font = f(size=8, color="666666"); r += 2
    # 서명란
    ws.merge_cells(f"A{r}:D{r}"); ws.merge_cells(f"E{r}:H{r}")
    ws[f"A{r}"] = "발  주  자"; ws[f"E{r}"] = "시  공  자"
    for col in ("A", "E"):
        ws[f"{col}{r}"].font = f(bold=True); ws[f"{col}{r}"].fill = G2; ws[f"{col}{r}"].alignment = C
    for j in range(1, 9):
        ws.cell(r, j).border = B_ALL
    r += 1
    ws.merge_cells(f"A{r}:D{r+2}"); ws.merge_cells(f"E{r}:H{r+2}")
    ws[f"A{r}"] = f"상  호 :  {proj.get('client','')}\n\n성  명 :  ______________________  (인)"
    ws[f"E{r}"] = f"상  호 :  {COMPANY['name']}\n\n성  명 :  {COMPANY['ceo']}                (인)"
    for col in ("A", "E"):
        ws[f"{col}{r}"].font = f(); ws[f"{col}{r}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    for rr in range(r, r + 3):
        for j in range(1, 9):
            ws.cell(rr, j).border = B_ALL
        ws.row_dimensions[rr].height = 22
    stamp = transparent_asset(find_asset(STAMP_FILE, assets_dir))
    if stamp:
        img = XLImage(stamp); img.width = 62; img.height = 64
        ws.add_image(img, f"G{r}")
    r += 4
    # 부대조건
    ws[f"A{r}"] = "부대조건 및 특기사항"; ws[f"A{r}"].font = f(size=10, bold=True, color="1B2A41"); r += 1
    for n, cnd in enumerate(data.get("conditions") or DEFAULT_CONDITIONS, 1):
        ws.merge_cells(f"A{r}:H{r}"); ws[f"A{r}"] = f"{n}. {cnd}"; ws[f"A{r}"].font = f(size=8.5); ws[f"A{r}"].alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 26; r += 1
    r += 1
    ws.merge_cells(f"A{r}:H{r}"); ws[f"A{r}"] = f"입금계좌 : {COMPANY['bank']}"; ws[f"A{r}"].font = f(bold=True); ws[f"A{r}"].fill = ORL; ws[f"A{r}"].alignment = C
    ws.row_dimensions[r].height = 20; r += 2
    ws.merge_cells(f"A{r}:H{r}"); ws[f"A{r}"] = f"상기와 같이 견적서를 제출합니다.      {COMPANY['name']}   대표   {COMPANY['ceo']}   (인)"
    ws[f"A{r}"].font = f(size=10, bold=True); ws[f"A{r}"].alignment = C; ws.row_dimensions[r].height = 24
    ws.page_setup.orientation = "portrait"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False
    wb.save(out_path)
    # 수식 캐시 계산 (LibreOffice 있으면)
    recalc = "/mnt/skills/public/xlsx/scripts/recalc.py"
    if os.path.isfile(recalc):
        try:
            subprocess.run([sys.executable, recalc, out_path, "60"], capture_output=True, timeout=120)
        except Exception:
            pass


# ───────────────────────────── main ─────────────────────────────
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("--out", default="/mnt/user-data/outputs")
    ap.add_argument("--assets", default=None)
    ap.add_argument("--check", action="store_true", help="계산만 하고 파일은 만들지 않음")
    a = ap.parse_args()
    with open(a.input, encoding="utf-8") as fp:
        data = json.load(fp)
    res = compute(data)
    summary = {"mode": res["mode"], "direct_mat": res["direct_mat"], "direct_lab": res["direct_lab"], "direct": res["direct"], "scale": res["scale"],
               "labor_ratio": round(res["direct_lab"] / res["direct"], 3) if res["direct"] else 0,
               "indirect": [(n, v) for n, s, v, nt in res["indirect"]] if res["mode"] == "lotte" else [],
               "total": res["total"], "vat": res["vat"], "grand": res["grand"], "total_kor": res["total_kor"], "fit": res["fit"]}
    print(json.dumps(summary, ensure_ascii=False, indent=1))
    if a.check:
        return
    os.makedirs(a.out, exist_ok=True)
    proj = data["project"]
    date_str = proj.get("date") or datetime.date.today().isoformat()
    safe = re.sub(r'[\\/:*?"<>|\[\]\s]+', "_", proj.get("site_name", "견적")).strip("_")[:40]
    base = os.path.join(a.out, f"유니온원_견적서_{safe}_{date_str}")
    build_pdf(data, res, base + ".pdf", a.assets)
    build_xlsx(data, res, base + ".xlsx", a.assets)
    print("OUTPUT:", base + ".pdf")
    print("OUTPUT:", base + ".xlsx")


if __name__ == "__main__":
    main()
